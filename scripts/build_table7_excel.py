#!/usr/bin/env python3
"""Build TABLE VII from existing in-vitro evaluation artifacts."""

from __future__ import annotations

import argparse
import json
import statistics
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path("/home/yangx/code")
PROJECT = ROOT / "new_deform/ai_worker/RTORv2"
PDF_PATH = ROOT / "new_deform/ai_worker/Resolving_the_Ambiguity_of_Complete-to-Partial_Point_Cloud_Registration_for_Image-Guided_Liver_Surgery_With_Patches-to-Partial_Matching——2026.pdf"
RTORV8 = ROOT / "new_deform/ai_worker/RTORv8/output/geotransformer.p2p_liver.rtor_a3_cooperative_abl1_no_proposal_seed7351/evaluation_epoch150"
GEO = ROOT / "new_deform/GeoTransformer_base/output/geotransformer.p2p_liver/evaluation_epoch150_standard"
CAST = ROOT / "Luo/CASTv2/output/p2p_liver"
DFAT = ROOT / "Luo/DFAT-main/output/geotransformer.p2p_liver"
PARE = ROOT / "point_clould/PARENet/output/P2P/evaluation_epoch150_rtorv2_standard"
LOW_VIS = PROJECT / "results/visibility_0.2_0.4_benchmark_metrics_v2"
ARCHIVE = ROOT / "point_clould/LiverMatch/Eva_in_vitro/Results_in_vitro/Result"

BIN_DEFINITIONS = (
    ("[0.2, 0.3)", 0.2, 0.3, False),
    ("[0.3, 0.4)", 0.3, 0.4, False),
    ("[0.4, 0.5)", 0.4, 0.5, False),
    ("[0.5, 0.6)", 0.5, 0.6, False),
    ("[0.6, 0.7)", 0.6, 0.7, False),
    ("[0.7, 0.8)", 0.7, 0.8, False),
    ("[0.8, 0.9)", 0.8, 0.9, False),
    ("[0.9, 1.0]", 0.9, 1.0, True),
)
BIN_LABELS = tuple(item[0] for item in BIN_DEFINITIONS)

METHODS = [
    {
        "name": "Ours", "env": "geo_v2", "source_type": "json",
        "requested": ROOT / "new_deform/ai_worker/RTORv8/output/geotransformer.p2p_liver.rtor_a3_cooperative_abl1_no_proposal_seed7351",
        "result": RTORV8 / "in_vitro_noise_none.json",
        "validation": "结果文件内 checkpoint 指向 RTORv5 同名实验，需复核迁移历史",
    },
    {
        "name": "GeoTransformer", "env": "geotransformer", "source_type": "json",
        "requested": ROOT / "new_deform/GeoTransformer_base/output/geotransformer.p2p_liver/snapshots/epoch-150.pth.tar",
        "result": GEO / "in_vitro_noise_none.json", "validation": "当前完整 JSON；已核对",
    },
    {
        "name": "CASTV2", "env": "overpredator_py310", "source_type": "json",
        "requested": ROOT / "Luo/CASTv2/ckpt/p2p_liver/castv2-p2p-liver-epoch-150.pth.tar",
        "result": CAST / "castv2_epoch150_in_vitro_all_metrics.json", "validation": "当前完整 JSON；已核对",
    },
    {
        "name": "DFAT", "env": "dfat_both", "source_type": "json",
        "requested": ROOT / "Luo/DFAT-main/output/geotransformer.p2p_liver/snapshots/epoch-150.pth.tar",
        "result": DFAT / "test_in_vitro_epoch150_summary.json", "validation": "当前完整 JSON；已核对",
    },
    {
        "name": "Lepard", "env": "livermatch5090", "source_type": "npy",
        "requested": ROOT / "new_deform/Lepard/snapshot/p2p_liver-paper/p2p_liver_from_scratch/checkpoints/model_best_loss.pth",
        "result": ARCHIVE / "Lepard/0.npy", "visibility_source": RTORV8 / "in_vitro_noise_none.json",
        "validation": "论文归档 NPY；未内嵌 checkpoint；当前 RANSAC 低可见率复跑见对照页",
    },
    {
        "name": "Lepard+P2P", "env": "livermatch5090", "source_type": "npy",
        "requested": ROOT / "new_deform/Lepard/snapshot/p2p_liver-paper/p2p_liver_from_scratch/checkpoints/model_best_loss.pth",
        "result": ARCHIVE / "P2P/5_Lepard_cl.npy", "visibility_source": RTORV8 / "in_vitro_noise_none.json",
        "validation": "论文归档 K=5 P2P；并非当前简化实现；当前复跑见对照页",
    },
    {
        "name": "PARENet", "env": "pare5090", "source_type": "json",
        "requested": ROOT / "point_clould/PARENet/output/P2P/snapshots/20260919-194858/epoch-150.pth.tar",
        "result": PARE / "in_vitro.json", "validation": "当前完整 JSON；已核对",
    },
    {
        "name": "LiverMatch", "env": "livermatch5090", "source_type": "npy",
        "requested": ROOT / "point_clould/LiverMatch/snapshot/liver_new_task3_004_002/checkpoints/model_best_loss.pth",
        "result": ARCHIVE / "LiverMatch/0.npy", "visibility_source": RTORV8 / "in_vitro_noise_none.json",
        "validation": "论文归档 NPY；未内嵌 checkpoint；当前低可见率复跑见对照页",
    },
    {
        "name": "LiverMatch+P2P", "env": "livermatch5090", "source_type": "npy",
        "requested": ROOT / "point_clould/LiverMatch/snapshot/liver_new_task3_004_002/checkpoints/model_best_loss.pth",
        "result": ARCHIVE / "P2P/5_LiverMatch_cl.npy", "visibility_source": RTORV8 / "in_vitro_noise_none.json",
        "validation": "论文归档 K=5 P2P NPY；当前低可见率复跑见对照页",
    },
]

LOW_VISIBILITY_RERUNS = [
    {"name": "Lepard", "result": LOW_VIS / "lepard/in_vitro.json", "note": "当前 RANSAC 复跑；部分样本无结果"},
    {"name": "Lepard+P2P", "result": LOW_VIS / "lepard_p2p/in_vitro.json", "note": "当前简化 P2P 实现，待复核；部分样本无结果"},
    {"name": "LiverMatch", "result": LOW_VIS / "livermatch/in_vitro.json", "note": "当前低可见率复跑"},
    {"name": "LiverMatch+P2P", "result": LOW_VIS / "livermatch_p2p/in_vitro.json", "note": "当前低可见率复跑"},
]

PAPER_TABLE = {
    "Procrustes": [(3.66, 1.69), (3.67, 1.78), (3.91, 1.74), (3.58, 1.80), (3.43, 1.68), (3.62, 1.70), (3.77, 1.60), (3.65, 1.91)],
    "GO-ICP [19]": [(87.15, 25.91), (82.58, 25.45), (75.34, 25.53), (72.20, 30.03), (61.89, 30.91), (47.04, 36.66), (21.16, 26.16), (9.31, 19.08)],
    "RoITr [28]": [(33.80, 24.06), (17.03, 17.20), (12.67, 14.43), (8.34, 9.33), (5.70, 3.69), (4.66, 2.21), (4.39, 1.80), (4.10, 2.09)],
    "RegTr [26]": [(66.25, 16.15), (55.91, 19.19), (45.02, 20.19), (37.39, 21.67), (29.74, 20.48), (16.15, 13.07), (8.76, 2.66), (5.16, 1.80)],
    "LiverMatch [22]": [(20.54, 22.33), (12.46, 12.87), (8.26, 5.34), (6.69, 4.29), (5.26, 3.16), (4.66, 2.58), (4.18, 1.70), (3.85, 1.85)],
    "LiverMatch [22] + proposed P2P": [(14.97, 18.73), (9.21, 7.95), (6.98, 4.51), (5.66, 2.80), (4.74, 3.03), (4.50, 2.26), (4.17, 1.74), (3.83, 1.92)],
    "Lepard [13]": [(20.71, 23.39), (10.32, 12.03), (6.58, 4.33), (5.30, 2.62), (4.65, 3.09), (4.24, 1.96), (4.14, 1.72), (3.86, 1.96)],
    "Lepard [13] + proposed P2P": [(12.45, 15.72), (8.43, 9.64), (6.07, 3.12), (5.08, 2.58), (4.45, 2.38), (4.26, 2.01), (4.19, 1.78), (3.91, 2.02)],
}


def _metric_value(sample: dict) -> float | None:
    value = sample.get("rms_tre_mm", sample.get("RMS_TRE_mm"))
    return None if value is None else float(value)


def summarize_visibility_bins(samples: list[dict], definitions=BIN_DEFINITIONS) -> dict[str, dict]:
    summaries = {}
    for label, lower, upper, include_upper in definitions:
        values = []
        for sample in samples:
            visibility = float(sample["visibility"])
            in_bin = lower <= visibility <= upper if include_upper else lower <= visibility < upper
            value = _metric_value(sample)
            if in_bin and value is not None:
                values.append(value)
        if not values:
            raise ValueError(f"No valid samples in visibility bin {label}")
        summaries[label] = {
            "count": len(values),
            "mean_rms_tre_mm": statistics.fmean(values),
            "std_rms_tre_mm": statistics.pstdev(values),
        }
    return summaries


def load_json_summary(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)
    return {
        "bins": summarize_visibility_bins(payload["samples"]),
        "source": str(path),
        "visibility_source": "embedded in JSON",
        "checkpoint": payload.get("checkpoint", ""),
        "evaluated_samples": payload.get("evaluated_samples", len(payload["samples"])),
        "source_kind": "当前 JSON",
    }


def load_npy_summary(path: Path, visibility_source: Path) -> dict:
    values = np.load(path, allow_pickle=False).astype(np.float64, copy=False).reshape(-1)
    with visibility_source.open("r", encoding="utf-8") as handle:
        metadata = json.load(handle)
    visibility = [float(sample["visibility"]) for sample in metadata["samples"]]
    if len(values) != len(visibility):
        raise ValueError(f"NPY/value length mismatch: {len(values)} != {len(visibility)}")
    samples = [
        {"visibility": vis, "rms_tre_mm": float(value)}
        for vis, value in zip(visibility, values)
    ]
    return {
        "bins": summarize_visibility_bins(samples),
        "source": str(path),
        "visibility_source": str(visibility_source),
        "checkpoint": "",
        "evaluated_samples": len(values),
        "source_kind": "论文归档 NPY",
    }


def load_method(method: dict) -> dict:
    if method["source_type"] == "json":
        return load_json_summary(Path(method["result"]))
    if method["source_type"] == "npy":
        return load_npy_summary(Path(method["result"]), Path(method["visibility_source"]))
    raise ValueError(f"Unsupported source type: {method['source_type']}")


def style_header(ws, row: int, start: int, end: int, color: str = "1F4E78") -> None:
    fill = PatternFill("solid", fgColor=color)
    for col in range(start, end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_grid(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    side = Side(style="thin", color="B7B7B7")
    border = Border(left=side, right=side, top=side, bottom=side)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def build_workbook(output: Path, methods: list[dict] | None = None, low_visibility_reruns: list[dict] | None = None) -> None:
    methods = METHODS if methods is None else methods
    low_visibility_reruns = LOW_VISIBILITY_RERUNS if low_visibility_reruns is None else low_visibility_reruns
    extracted = [{**method, "summary": load_method(method)} for method in methods]

    wb = Workbook()
    ws = wb.active
    ws.title = "TABLE7_当前汇总"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    ws["A1"] = "TABLE VII — in vitro phantom 不同可见率区间的 RMS-TRE"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=10)
    ws["A2"] = "单位 mm；mean ± population std；当前完整 JSON 与论文归档 NPY 混合汇总，来源状态见最后一列"
    ws["A2"].alignment = Alignment(horizontal="center")
    headers = ["模型"] + list(BIN_LABELS) + ["来源与验证状态"]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    style_header(ws, 4, 1, len(headers))
    warning_fill = PatternFill("solid", fgColor="FCE4D6")
    complete_fill = PatternFill("solid", fgColor="E2F0D9")
    for row_index, method in enumerate(extracted, 5):
        ws.cell(row_index, 1, method["name"])
        for col, label in enumerate(BIN_LABELS, 2):
            item = method["summary"]["bins"][label]
            ws.cell(row_index, col, f"{item['mean_rms_tre_mm']:.2f} ± {item['std_rms_tre_mm']:.2f}")
        status = method.get("validation", "已核对")
        ws.cell(row_index, 10, status)
        ws.cell(row_index, 10).fill = warning_fill if ("归档" in status or "复核" in status) else complete_fill
    style_grid(ws, 4, 4 + len(extracted), 1, 10)
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:J{4 + len(extracted)}"
    ws.column_dimensions["A"].width = 22
    for col in range(2, 10):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions["J"].width = 70

    detail = wb.create_sheet("数值明细")
    headers = ["模型", "Visibility bin", "Mean RMS-TRE (mm)", "Std RMS-TRE (mm)", "N", "来源类型", "结果文件"]
    detail.append(headers)
    style_header(detail, 1, 1, len(headers))
    for method in extracted:
        for label in BIN_LABELS:
            item = method["summary"]["bins"][label]
            detail.append([method["name"], label, item["mean_rms_tre_mm"], item["std_rms_tre_mm"], item["count"], method["summary"]["source_kind"], method["summary"]["source"]])
    style_grid(detail, 1, detail.max_row, 1, 7)
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = f"A1:G{detail.max_row}"
    for row in range(2, detail.max_row + 1):
        detail.cell(row, 3).number_format = "0.0000"
        detail.cell(row, 4).number_format = "0.0000"
    for col, width in enumerate([22, 20, 22, 22, 10, 22, 115], 1):
        detail.column_dimensions[get_column_letter(col)].width = width

    provenance = wb.create_sheet("数据来源")
    headers = ["模型", "Conda 环境", "来源类型", "用户指定模型路径", "实际结果文件", "可见率来源", "结果内 checkpoint", "总样本数", "核对说明"]
    provenance.append(headers)
    style_header(provenance, 1, 1, len(headers))
    for method in extracted:
        summary = method["summary"]
        provenance.append([method["name"], method["env"], summary["source_kind"], str(method["requested"]), summary["source"], summary["visibility_source"], summary["checkpoint"], summary["evaluated_samples"], method.get("validation", "已核对")])
    style_grid(provenance, 1, provenance.max_row, 1, 9)
    provenance.freeze_panes = "A2"
    for col, width in enumerate([22, 20, 22, 85, 115, 115, 90, 15, 75], 1):
        provenance.column_dimensions[get_column_letter(col)].width = width

    rerun_ws = wb.create_sheet("当前低可见率复跑")
    rerun_headers = ["模型", "当前 [0.2,0.3)", "归档 [0.2,0.3)", "当前-归档均值", "当前 N", "当前 [0.3,0.4)", "归档 [0.3,0.4)", "当前-归档均值", "当前 N", "说明", "当前结果文件"]
    rerun_ws.append(rerun_headers)
    style_header(rerun_ws, 1, 1, len(rerun_headers), color="C65911")
    archive_lookup = {method["name"]: method["summary"]["bins"] for method in extracted}
    for item in low_visibility_reruns:
        with Path(item["result"]).open("r", encoding="utf-8") as handle:
            rerun_payload = json.load(handle)
        current = summarize_visibility_bins(rerun_payload["samples"], BIN_DEFINITIONS[:2])
        archive = archive_lookup[item["name"]]
        row = [item["name"]]
        for label in ("[0.2, 0.3)", "[0.3, 0.4)"):
            cur, arc = current[label], archive[label]
            row.extend([
                f"{cur['mean_rms_tre_mm']:.2f} ± {cur['std_rms_tre_mm']:.2f}",
                f"{arc['mean_rms_tre_mm']:.2f} ± {arc['std_rms_tre_mm']:.2f}",
                cur["mean_rms_tre_mm"] - arc["mean_rms_tre_mm"], cur["count"],
            ])
        row.extend([item["note"], str(item["result"])])
        rerun_ws.append(row)
    style_grid(rerun_ws, 1, rerun_ws.max_row, 1, 11)
    for row in range(2, rerun_ws.max_row + 1):
        rerun_ws.cell(row, 4).number_format = "+0.0000;-0.0000;0.0000"
        rerun_ws.cell(row, 8).number_format = "+0.0000;-0.0000;0.0000"
        for col in range(1, 12):
            rerun_ws.cell(row, col).fill = warning_fill
    for col, width in enumerate([22, 22, 22, 20, 12, 22, 22, 20, 12, 55, 115], 1):
        rerun_ws.column_dimensions[get_column_letter(col)].width = width

    paper = wb.create_sheet("PDF原TABLE7")
    paper.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    paper["A1"] = "PDF TABLE VII 原始数值 — RMS-TRE（mm，mean ± std）"
    paper["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    paper["A1"].fill = PatternFill("solid", fgColor="7030A0")
    paper["A1"].alignment = Alignment(horizontal="center")
    paper.merge_cells(start_row=2, start_column=1, end_row=2, end_column=9)
    paper["A2"] = f"来源：{PDF_PATH}；印刷页 468（PDF 第 10 页）；仅作论文原值对照"
    paper["A2"].alignment = Alignment(horizontal="center", wrap_text=True)
    for col, value in enumerate(["Method"] + list(BIN_LABELS), 1):
        paper.cell(4, col, value)
    style_header(paper, 4, 1, 9, color="7030A0")
    for method_name, values in PAPER_TABLE.items():
        paper.append([method_name] + [f"{mean:.2f} ± {std:.2f}" for mean, std in values])
    style_grid(paper, 4, paper.max_row, 1, 9)
    paper.column_dimensions["A"].width = 42
    for col in range(2, 10):
        paper.column_dimensions[get_column_letter(col)].width = 18

    readme = wb.create_sheet("说明")
    notes = [
        ["项目", "内容"],
        ["验证状态", "ANALYZED：只聚合已有 JSON/NPY 结果，没有重新运行模型。"],
        ["指标", "RMS-TRE，单位 mm；均值 ± 总体标准差（ddof=0）。"],
        ["分箱", "前七档采用左闭右开；最后一档为 [0.9,1.0]，包含 visibility=1.0；范围外样本不纳入表格。"],
        ["主表来源", "Ours、GeoTransformer、CASTV2、DFAT、PARENet 使用当前完整 JSON；Lepard/LiverMatch 四行使用原仓库论文归档 NPY。"],
        ["归档 NPY 连接", "归档 NPY 不含样本名或可见率；按原 Table_VII.py 的位置索引逻辑，与同顺序的 800 例 in vitro 可见率元数据连接。所得四行逐格复现 PDF TABLE VII。"],
        ["checkpoint 限制", "归档 NPY 不内嵌 checkpoint，文件夹与模型的对应来自原仓库脚本；不能视为对用户指定 checkpoint 的重新验证。"],
        ["低可见率对照", "当前低可见率复跑单列一页。Lepard 当前复跑仅有 204 条记录，其中两个目标分箱分别有 100/95 条，少于完整数据的 105/99 条。"],
        ["Lepard+P2P", "主表使用论文归档 K=5 P2P；当前简化 P2P 复跑不是同一实现，结果不得混用。"],
        ["Ours 路径", "结果位于 RTORv8，但 JSON 内 checkpoint 指向 RTORv5 同名实验，保留迁移历史复核提示。"],
        ["统计谬误扫描", "11/11 已检查：本表为描述性分组汇总，不作显著性或因果推断；来源混合、样本缺失和多重比较均保持显式。"],
    ]
    for row in notes:
        readme.append(row)
    style_header(readme, 1, 1, 2, color="548235")
    style_grid(readme, 1, readme.max_row, 1, 2)
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 130

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    reopened = load_workbook(output, read_only=True)
    expected = ["TABLE7_当前汇总", "数值明细", "数据来源", "当前低可见率复跑", "PDF原TABLE7", "说明"]
    if reopened.sheetnames != expected:
        raise RuntimeError(f"Unexpected workbook sheets: {reopened.sheetnames}")
    if reopened["TABLE7_当前汇总"].max_row != 4 + len(extracted):
        raise RuntimeError("Unexpected number of method rows")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output", type=Path,
        default=PROJECT / "results/paper_tables/TABLE7_in_vitro_visibility_RMS_TRE_20260924.xlsx",
    )
    args = parser.parse_args()
    build_workbook(args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
