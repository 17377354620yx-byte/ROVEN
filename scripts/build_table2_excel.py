#!/usr/bin/env python3
"""Build the paper TABLE II workbook from existing evaluation results."""

from __future__ import annotations

import argparse
import json
import statistics
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path("/home/yangx/code")
PROJECT = ROOT / "new_deform/ai_worker/RTORv2"
PDF_PATH = ROOT / "new_deform/ai_worker/Resolving_the_Ambiguity_of_Complete-to-Partial_Point_Cloud_Registration_for_Image-Guided_Liver_Surgery_With_Patches-to-Partial_Matching——2026.pdf"
NOISE_LEVELS = (0, 2, 4)

RTORV8 = ROOT / "new_deform/ai_worker/RTORv8/output/geotransformer.p2p_liver.rtor_a3_cooperative_abl1_no_proposal_seed7351/evaluation_epoch150"
GEO = ROOT / "new_deform/GeoTransformer_base/output/geotransformer.p2p_liver/evaluation_epoch150_standard"
CAST = ROOT / "Luo/CASTv2/output/p2p_liver"
DFAT = ROOT / "Luo/DFAT-main/output/geotransformer.p2p_liver"
PARE = ROOT / "point_clould/PARENet/output/P2P/evaluation_epoch150_rtorv2_standard"
LOW_VIS = PROJECT / "results/visibility_0.2_0.4_benchmark_metrics_v2"

METHODS = [
    {
        "name": "Ours",
        "env": "geo_v2",
        "requested": ROOT / "new_deform/ai_worker/RTORv8/output/geotransformer.p2p_liver.rtor_a3_cooperative_abl1_no_proposal_seed7351",
        "results": {
            0: RTORV8 / "in_silico_noise_none.json",
            2: RTORV8 / "in_silico_noise_2.json",
            4: RTORV8 / "in_silico_noise_4.json",
        },
        "validation": "结果文件内 checkpoint 字段指向 RTORv5 同名实验，需复核迁移历史",
    },
    {
        "name": "GeoTransformer",
        "env": "geotransformer",
        "requested": ROOT / "new_deform/GeoTransformer_base/output/geotransformer.p2p_liver/snapshots/epoch-150.pth.tar",
        "results": {
            0: GEO / "in_silico_noise_none.json",
            2: GEO / "in_silico_noise_2.json",
            4: GEO / "in_silico_noise_4.json",
        },
        "validation": "已核对",
    },
    {
        "name": "CASTV2",
        "env": "overpredator_py310",
        "requested": ROOT / "Luo/CASTv2/ckpt/p2p_liver/castv2-p2p-liver-epoch-150.pth.tar",
        "results": {
            0: CAST / "castv2_epoch150_in_silico_all_metrics.json",
            2: None,
            4: LOW_VIS / "castv2/in_silico_noise4.json",
        },
        "validation": "缺少 2 mm 结果",
    },
    {
        "name": "DFAT",
        "env": "dfat_both",
        "requested": ROOT / "Luo/DFAT-main/output/geotransformer.p2p_liver/snapshots/epoch-150.pth.tar",
        "results": {
            0: DFAT / "test_in_silico_epoch150_tableVI.json",
            2: None,
            4: LOW_VIS / "dfat/in_silico_noise4.json",
        },
        "validation": "缺少 2 mm 结果",
    },
    {
        "name": "Lepard",
        "env": "livermatch5090",
        "requested": ROOT / "new_deform/Lepard/snapshot/p2p_liver-paper/p2p_liver_from_scratch/checkpoints/model_best_loss.pth",
        "results": {
            0: LOW_VIS / "lepard/in_silico_none.json",
            2: None,
            4: LOW_VIS / "lepard/in_silico_noise4.json",
        },
        "validation": "RANSAC 基线，论文口径待复核；缺少 2 mm 结果",
    },
    {
        "name": "Lepard+P2P",
        "env": "livermatch5090",
        "requested": ROOT / "new_deform/Lepard/snapshot/p2p_liver-paper/p2p_liver_from_scratch/checkpoints/model_best_loss.pth",
        "results": {
            0: LOW_VIS / "lepard_p2p/in_silico_none.json",
            2: None,
            4: LOW_VIS / "lepard_p2p/in_silico_noise4.json",
        },
        "validation": "简化 P2P 实现，待复核；缺少 2 mm 结果",
    },
    {
        "name": "PARENet",
        "env": "pare5090",
        "requested": ROOT / "point_clould/PARENet/output/P2P/snapshots/20260919-194858/epoch-150.pth.tar",
        "results": {
            0: PARE / "in_silico_noise_none.json",
            2: PARE / "in_silico_noise_2.json",
            4: PARE / "in_silico_noise_4.json",
        },
        "validation": "已核对",
    },
    {
        "name": "LiverMatch",
        "env": "livermatch5090",
        "requested": ROOT / "point_clould/LiverMatch/snapshot/liver_new_task3_004_002/checkpoints/model_best_loss.pth",
        "results": {
            0: LOW_VIS / "livermatch/in_silico_none.json",
            2: None,
            4: LOW_VIS / "livermatch/in_silico_noise4.json",
        },
        "validation": "缺少 2 mm 结果",
    },
    {
        "name": "LiverMatch+P2P",
        "env": "livermatch5090",
        "requested": ROOT / "point_clould/LiverMatch/snapshot/liver_new_task3_004_002/checkpoints/model_best_loss.pth",
        "results": {
            0: LOW_VIS / "livermatch_p2p/in_silico_none.json",
            2: None,
            4: LOW_VIS / "livermatch_p2p/in_silico_noise4.json",
        },
        "validation": "缺少 2 mm 结果",
    },
]

PAPER_TABLE = {
    "Procrustes": [(3.37, 2.11), (3.37, 2.11), (3.37, 2.11)],
    "LiverMatch [22]": [(12.63, 13.87), (13.72, 15.48), (21.85, 23.20)],
    "LiverMatch [22] + proposed P2P": [(8.74, 10.23), (9.19, 9.48), (14.58, 18.11)],
    "Lepard [13]": [(9.47, 14.74), (10.07, 15.28), (14.70, 21.02)],
    "Lepard [13] + proposed P2P": [(6.73, 5.96), (7.22, 8.45), (9.85, 14.88)],
}


def _metric_value(sample: dict) -> float | None:
    value = sample.get("rms_tre_mm", sample.get("RMS_TRE_mm"))
    return None if value is None else float(value)


def summarize_low_visibility(samples: list[dict]) -> dict:
    values = [
        value
        for sample in samples
        if 0.2 <= float(sample["visibility"]) < 0.3
        if (value := _metric_value(sample)) is not None
    ]
    if not values:
        raise ValueError("No valid samples in visibility range [0.2, 0.3)")
    return {
        "count": len(values),
        "mean_rms_tre_mm": statistics.fmean(values),
        "std_rms_tre_mm": statistics.pstdev(values),
    }


def load_summary(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)
    summary = summarize_low_visibility(payload["samples"])
    summary.update(
        {
            "source": str(path),
            "checkpoint": payload.get("checkpoint", ""),
            "evaluated_samples": payload.get("evaluated_samples", len(payload["samples"])),
        }
    )
    return summary


def style_header(ws, row: int, start: int, end: int, color: str = "1F4E78") -> None:
    fill = PatternFill("solid", fgColor=color)
    for col in range(start, end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_grid(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    side = Side(style="thin", color="B7B7B7")
    border = Border(left=side, right=side, top=side, bottom=side)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def build_workbook(output: Path, methods: list[dict] | None = None) -> None:
    methods = METHODS if methods is None else methods
    extracted = []
    for method in methods:
        summaries = {}
        for noise in NOISE_LEVELS:
            path = method["results"].get(noise)
            summaries[noise] = None if path is None else load_summary(Path(path))
        extracted.append({**method, "summaries": summaries})

    wb = Workbook()
    ws = wb.active
    ws.title = "TABLE2_当前结果"
    ws.merge_cells("A1:E1")
    ws["A1"] = "TABLE II — 低可见度区间不同噪声水平的 RMS-TRE"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:E2")
    ws["A2"] = "in silico phantom；visibility ∈ [0.2,0.3)；单位 mm；mean ± std"
    ws["A2"].alignment = Alignment(horizontal="center")
    headers = ["模型", "0 mm", "2 mm", "4 mm", "覆盖与验证状态"]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    style_header(ws, 4, 1, 5)
    missing_fill = PatternFill("solid", fgColor="FFF2CC")
    warning_fill = PatternFill("solid", fgColor="FCE4D6")
    complete_fill = PatternFill("solid", fgColor="E2F0D9")
    for row_index, method in enumerate(extracted, 5):
        ws.cell(row_index, 1, method["name"])
        available = 0
        for col, noise in enumerate(NOISE_LEVELS, 2):
            item = method["summaries"][noise]
            cell = ws.cell(row_index, col)
            if item is None:
                cell.value = "未找到结果"
                cell.fill = missing_fill
            else:
                available += 1
                cell.value = f"{item['mean_rms_tre_mm']:.2f} ± {item['std_rms_tre_mm']:.2f}"
        validation = method.get("validation", "已核对")
        ws.cell(row_index, 5, f"{available}/3；{validation}")
        if "待复核" in validation or "需复核" in validation:
            ws.cell(row_index, 5).fill = warning_fill
        elif available == 3:
            ws.cell(row_index, 5).fill = complete_fill
        else:
            ws.cell(row_index, 5).fill = missing_fill
    style_grid(ws, 4, 4 + len(extracted), 1, 5)
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:E{4 + len(extracted)}"
    for col, width in enumerate([24, 20, 20, 20, 54], 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    detail = wb.create_sheet("数值明细")
    detail_headers = ["模型", "Noise (mm)", "Mean RMS-TRE (mm)", "Std RMS-TRE (mm)", "N ([0.2,0.3))", "状态", "结果文件"]
    detail.append(detail_headers)
    style_header(detail, 1, 1, len(detail_headers))
    for method in extracted:
        for noise in NOISE_LEVELS:
            item = method["summaries"][noise]
            detail.append(
                [
                    method["name"], noise,
                    None if item is None else item["mean_rms_tre_mm"],
                    None if item is None else item["std_rms_tre_mm"],
                    None if item is None else item["count"],
                    "缺失" if item is None else method.get("validation", "已核对"),
                    "" if item is None else item["source"],
                ]
            )
    style_grid(detail, 1, detail.max_row, 1, 7)
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = f"A1:G{detail.max_row}"
    for row in range(2, detail.max_row + 1):
        detail.cell(row, 3).number_format = "0.0000"
        detail.cell(row, 4).number_format = "0.0000"
        if detail.cell(row, 6).value == "缺失":
            for col in range(1, 8):
                detail.cell(row, col).fill = missing_fill
    for col, width in enumerate([22, 14, 22, 22, 18, 48, 105], 1):
        detail.column_dimensions[get_column_letter(col)].width = width

    provenance = wb.create_sheet("数据来源")
    provenance_headers = ["模型", "Noise (mm)", "Conda 环境", "用户指定模型路径", "实际结果文件", "结果内 checkpoint", "评测总样本数", "区间样本数", "核对说明"]
    provenance.append(provenance_headers)
    style_header(provenance, 1, 1, len(provenance_headers))
    for method in extracted:
        for noise in NOISE_LEVELS:
            item = method["summaries"][noise]
            provenance.append(
                [
                    method["name"], noise, method["env"], str(method["requested"]),
                    "" if item is None else item["source"],
                    "" if item is None else item["checkpoint"],
                    None if item is None else item["evaluated_samples"],
                    None if item is None else item["count"],
                    "未找到现有结果" if item is None else method.get("validation", "已核对"),
                ]
            )
    style_grid(provenance, 1, provenance.max_row, 1, 9)
    provenance.freeze_panes = "A2"
    for col, width in enumerate([22, 13, 20, 82, 105, 82, 18, 16, 58], 1):
        provenance.column_dimensions[get_column_letter(col)].width = width

    paper = wb.create_sheet("PDF原TABLE2")
    paper.merge_cells("A1:D1")
    paper["A1"] = "PDF TABLE II 原始数值 — RMS-TRE（mm，mean ± std）"
    paper["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    paper["A1"].fill = PatternFill("solid", fgColor="7030A0")
    paper["A1"].alignment = Alignment(horizontal="center")
    paper.merge_cells("A2:D2")
    paper["A2"] = f"来源：{PDF_PATH}；印刷页 466（PDF 第 8 页）；visibility ∈ [0.2,0.3)；仅作论文原值对照"
    paper["A2"].alignment = Alignment(horizontal="center", wrap_text=True)
    for col, value in enumerate(["Method", "0 mm", "2 mm", "4 mm"], 1):
        paper.cell(4, col, value)
    style_header(paper, 4, 1, 4, color="7030A0")
    for method_name, values in PAPER_TABLE.items():
        paper.append([method_name] + [f"{mean:.2f} ± {std:.2f}" for mean, std in values])
    style_grid(paper, 4, paper.max_row, 1, 4)
    paper.column_dimensions["A"].width = 40
    for col in range(2, 5):
        paper.column_dimensions[get_column_letter(col)].width = 20

    readme = wb.create_sheet("说明")
    notes = [
        ["项目", "内容"],
        ["指标", "RMS-TRE on paired volumetric fiducials，单位 mm，表中为均值 ± 总体标准差（ddof=0）。"],
        ["筛选", "严格采用 0.2 ≤ visibility < 0.3；visibility=0.3 的边界样本不纳入。"],
        ["噪声", "0、2、4 mm；0 mm 对应 noise=None。"],
        ["缺失处理", "没有现成结果的组合标记为“未找到结果”，不以 PDF 原论文数值填补当前 checkpoint。"],
        ["Lepard", "当前 Lepard 使用 RANSAC 基线；Lepard+P2P 是简化实现，两者均不宜直接作为论文最终对照。"],
        ["统计限制", "这里只进行描述性汇总，没有执行显著性检验；不同噪声水平应在同一测试样本集合上进行配对比较。"],
        ["验证状态", "ANALYZED：读取并汇总既有结果，未重新运行模型。统计谬误扫描 11/11；主要警示为缺失噪声组合、低可见度子集外推和 Lepard 求解器口径。"],
        ["生成日期", "2026-09-24（Asia/Shanghai）"],
    ]
    for row in notes:
        readme.append(row)
    style_header(readme, 1, 1, 2)
    style_grid(readme, 1, readme.max_row, 1, 2)
    readme.column_dimensions["A"].width = 22
    readme.column_dimensions["B"].width = 122
    for row in range(1, readme.max_row + 1):
        readme.row_dimensions[row].height = 36

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    checked = load_workbook(output, read_only=True, data_only=False)
    expected = ["TABLE2_当前结果", "数值明细", "数据来源", "PDF原TABLE2", "说明"]
    if checked.sheetnames != expected:
        raise RuntimeError(f"Unexpected workbook sheets: {checked.sheetnames}")
    if checked["TABLE2_当前结果"].max_row != 4 + len(methods):
        raise RuntimeError("Current-results sheet has an unexpected method count")
    checked.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("results/paper_tables/TABLE2_in_silico_noise_visibility_0.2_0.3_RMS_TRE_20260924.xlsx"),
    )
    args = parser.parse_args()
    build_workbook(args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
