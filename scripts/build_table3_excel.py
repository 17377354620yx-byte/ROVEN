#!/usr/bin/env python3
"""Build TABLE III from existing no-noise in-silico evaluation JSON files."""

from __future__ import annotations

import argparse
import json
import statistics
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path("/home/yangx/code")
PROJECT = ROOT / "new_deform/ai_worker/RTORv2"
PDF_PATH = ROOT / "new_deform/ai_worker/Resolving_the_Ambiguity_of_Complete-to-Partial_Point_Cloud_Registration_for_Image-Guided_Liver_Surgery_With_Patches-to-Partial_Matching——2026.pdf"
RTORV8 = ROOT / "new_deform/ai_worker/RTORv8/output/geotransformer.p2p_liver.rtor_a3_cooperative_abl1_no_proposal_seed7351/evaluation_epoch150"
GEO = ROOT / "new_deform/GeoTransformer_base/output/geotransformer.p2p_liver/evaluation_epoch150_standard"
CAST = ROOT / "Luo/CASTv2/output/p2p_liver"
DFAT = ROOT / "Luo/DFAT-main/output/geotransformer.p2p_liver"
PARE = ROOT / "point_clould/PARENet/output/P2P/evaluation_epoch150_rtorv2_standard"
LOW_VIS = PROJECT / "results/visibility_0.2_0.4_benchmark_metrics_v2"
DEFORMATION_RANGES = ("[0, 6)", "[6, 12]")

METHODS = [
    {
        "name": "Ours",
        "env": "geo_v2",
        "requested": ROOT / "new_deform/ai_worker/RTORv8/output/geotransformer.p2p_liver.rtor_a3_cooperative_abl1_no_proposal_seed7351",
        "result": RTORV8 / "in_silico_noise_none.json",
        "validation": "结果文件内 checkpoint 字段指向 RTORv5 同名实验，需复核迁移历史",
    },
    {
        "name": "GeoTransformer",
        "env": "geotransformer",
        "requested": ROOT / "new_deform/GeoTransformer_base/output/geotransformer.p2p_liver/snapshots/epoch-150.pth.tar",
        "result": GEO / "in_silico_noise_none.json",
        "validation": "已核对",
    },
    {
        "name": "CASTV2",
        "env": "overpredator_py310",
        "requested": ROOT / "Luo/CASTv2/ckpt/p2p_liver/castv2-p2p-liver-epoch-150.pth.tar",
        "result": CAST / "castv2_epoch150_in_silico_all_metrics.json",
        "validation": "已核对",
    },
    {
        "name": "DFAT",
        "env": "dfat_both",
        "requested": ROOT / "Luo/DFAT-main/output/geotransformer.p2p_liver/snapshots/epoch-150.pth.tar",
        "result": DFAT / "test_in_silico_epoch150_tableVI.json",
        "validation": "已核对",
    },
    {
        "name": "Lepard",
        "env": "livermatch5090",
        "requested": ROOT / "new_deform/Lepard/snapshot/p2p_liver-paper/p2p_liver_from_scratch/checkpoints/model_best_loss.pth",
        "result": LOW_VIS / "lepard/in_silico_none.json",
        "metadata": RTORV8 / "in_silico_noise_none.json",
        "validation": "RANSAC 基线，论文口径待复核",
    },
    {
        "name": "Lepard+P2P",
        "env": "livermatch5090",
        "requested": ROOT / "new_deform/Lepard/snapshot/p2p_liver-paper/p2p_liver_from_scratch/checkpoints/model_best_loss.pth",
        "result": LOW_VIS / "lepard_p2p/in_silico_none.json",
        "metadata": RTORV8 / "in_silico_noise_none.json",
        "validation": "简化 P2P 实现，待复核",
    },
    {
        "name": "PARENet",
        "env": "pare5090",
        "requested": ROOT / "point_clould/PARENet/output/P2P/snapshots/20260919-194858/epoch-150.pth.tar",
        "result": PARE / "in_silico_noise_none.json",
        "validation": "已核对",
    },
    {
        "name": "LiverMatch",
        "env": "livermatch5090",
        "requested": ROOT / "point_clould/LiverMatch/snapshot/liver_new_task3_004_002/checkpoints/model_best_loss.pth",
        "result": LOW_VIS / "livermatch/in_silico_none.json",
        "validation": "已核对",
    },
    {
        "name": "LiverMatch+P2P",
        "env": "livermatch5090",
        "requested": ROOT / "point_clould/LiverMatch/snapshot/liver_new_task3_004_002/checkpoints/model_best_loss.pth",
        "result": LOW_VIS / "livermatch_p2p/in_silico_none.json",
        "validation": "已核对",
    },
]

PAPER_TABLE = {
    "Procrustes": [(2.82, 1.51), (7.30, 1.25)],
    "LiverMatch [22]": [(12.54, 13.75), (17.24, 15.93)],
    "LiverMatch [22] + proposed P2P": [(8.54, 11.16), (12.54, 5.13)],
    "Lepard [13]": [(8.92, 14.66), (15.50, 13.82)],
    "Lepard [13] + proposed P2P": [(6.19, 5.50), (12.85, 7.24)],
}


def _metric_value(sample: dict) -> float | None:
    value = sample.get("rms_tre_mm", sample.get("RMS_TRE_mm"))
    return None if value is None else float(value)


def _sample_name(sample: dict) -> str | None:
    return sample.get("sample", sample.get("sample_name"))


def summarize_deformation_bins(samples: list[dict]) -> dict[str, dict]:
    definitions = (
        ("[0, 6)", 0.0, 6.0, False),
        ("[6, 12]", 6.0, 12.0, True),
    )
    summaries = {}
    for label, lower, upper, include_upper in definitions:
        values = []
        for sample in samples:
            visibility = float(sample["visibility"])
            deformation = float(sample["deformation_mm"])
            in_deformation_range = lower <= deformation <= upper if include_upper else lower <= deformation < upper
            value = _metric_value(sample)
            if 0.2 <= visibility < 0.3 and in_deformation_range and value is not None:
                values.append(value)
        if not values:
            raise ValueError(f"No valid samples for deformation range {label}")
        summaries[label] = {
            "count": len(values),
            "mean_rms_tre_mm": statistics.fmean(values),
            "std_rms_tre_mm": statistics.pstdev(values),
        }
    return summaries


def _add_deformation_metadata(samples: list[dict], metadata_path: Path) -> tuple[list[dict], int]:
    with metadata_path.open("r", encoding="utf-8") as handle:
        metadata_payload = json.load(handle)
    lookup = {
        _sample_name(sample): float(sample["deformation_mm"])
        for sample in metadata_payload["samples"]
        if _sample_name(sample) is not None and sample.get("deformation_mm") is not None
    }
    enriched = []
    joined = 0
    missing = []
    for sample in samples:
        row = dict(sample)
        if row.get("deformation_mm") is None:
            name = _sample_name(row)
            if name not in lookup:
                missing.append(name)
            else:
                row["deformation_mm"] = lookup[name]
                joined += 1
        enriched.append(row)
    if missing:
        raise ValueError(f"Missing deformation metadata for {len(missing)} samples; first={missing[0]!r}")
    return enriched, joined


def load_summary(path: Path, metadata_path: Path | None = None) -> dict:
    with path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)
    samples = payload["samples"]
    joined = 0
    if any(sample.get("deformation_mm") is None for sample in samples):
        if metadata_path is None:
            raise ValueError(f"{path} lacks deformation_mm and no metadata source was provided")
        samples, joined = _add_deformation_metadata(samples, metadata_path)
    summaries = summarize_deformation_bins(samples)
    return {
        "bins": summaries,
        "source": str(path),
        "metadata_source": "" if metadata_path is None else str(metadata_path),
        "metadata_joined": joined,
        "checkpoint": payload.get("checkpoint", ""),
        "evaluated_samples": payload.get("evaluated_samples", len(samples)),
    }


def style_header(ws, row: int, start: int, end: int, color: str = "1F4E78") -> None:
    fill = PatternFill("solid", fgColor=color)
    for col in range(start, end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_grid(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    side = Side(style="thin", color="B7B7B7")
    border = Border(left=side, right=side, top=side, bottom=side)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def _status(method: dict, summary: dict) -> str:
    validation = method.get("validation", "已核对")
    if "简化 P2P 实现" in validation:
        validation = "实现待复核"
    if summary["metadata_joined"]:
        validation += "；deformation 元数据由样本名匹配补充"
    return validation


def build_workbook(output: Path, methods: list[dict] | None = None) -> None:
    methods = METHODS if methods is None else methods
    extracted = []
    for method in methods:
        metadata_path = method.get("metadata")
        summary = load_summary(
            Path(method["result"]),
            None if metadata_path is None else Path(metadata_path),
        )
        extracted.append({**method, "summary": summary})

    wb = Workbook()
    ws = wb.active
    ws.title = "TABLE3_当前结果"
    ws.merge_cells("A1:D1")
    ws["A1"] = "TABLE III — 低可见度区间不同形变范围的 RMS-TRE"
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:D2")
    ws["A2"] = "in silico phantom；noise = 0 mm；visibility ∈ [0.2,0.3)；单位 mm；mean ± std"
    ws["A2"].alignment = Alignment(horizontal="center")
    for col, value in enumerate(["模型", "[0, 6) mm", "[6, 12] mm", "验证状态"], 1):
        ws.cell(4, col, value)
    style_header(ws, 4, 1, 4)
    warning_fill = PatternFill("solid", fgColor="FCE4D6")
    complete_fill = PatternFill("solid", fgColor="E2F0D9")
    for row_index, method in enumerate(extracted, 5):
        ws.cell(row_index, 1, method["name"])
        for col, label in enumerate(DEFORMATION_RANGES, 2):
            item = method["summary"]["bins"][label]
            ws.cell(row_index, col, f"{item['mean_rms_tre_mm']:.2f} ± {item['std_rms_tre_mm']:.2f}")
        status = _status(method, method["summary"])
        ws.cell(row_index, 4, status)
        ws.cell(row_index, 4).fill = warning_fill if "复核" in status else complete_fill
    style_grid(ws, 4, 4 + len(extracted), 1, 4)
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:D{4 + len(extracted)}"
    for col, width in enumerate([24, 22, 22, 60], 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    detail = wb.create_sheet("数值明细")
    headers = ["模型", "Deformation range (mm)", "Mean RMS-TRE (mm)", "Std RMS-TRE (mm)", "N", "状态", "结果文件"]
    detail.append(headers)
    style_header(detail, 1, 1, len(headers))
    for method in extracted:
        for label in DEFORMATION_RANGES:
            item = method["summary"]["bins"][label]
            detail.append([
                method["name"], label, item["mean_rms_tre_mm"], item["std_rms_tre_mm"], item["count"],
                _status(method, method["summary"]), method["summary"]["source"],
            ])
    style_grid(detail, 1, detail.max_row, 1, 7)
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = f"A1:G{detail.max_row}"
    for row in range(2, detail.max_row + 1):
        detail.cell(row, 3).number_format = "0.0000"
        detail.cell(row, 4).number_format = "0.0000"
    for col, width in enumerate([22, 25, 22, 22, 10, 60, 110], 1):
        detail.column_dimensions[get_column_letter(col)].width = width

    provenance = wb.create_sheet("数据来源")
    headers = ["模型", "Conda 环境", "用户指定模型路径", "实际结果文件", "deformation 元数据文件", "结果内 checkpoint", "评测总样本数", "元数据匹配数", "核对说明"]
    provenance.append(headers)
    style_header(provenance, 1, 1, len(headers))
    for method in extracted:
        summary = method["summary"]
        provenance.append([
            method["name"], method["env"], str(method["requested"]), summary["source"],
            summary["metadata_source"], summary["checkpoint"], summary["evaluated_samples"],
            summary["metadata_joined"], _status(method, summary),
        ])
    style_grid(provenance, 1, provenance.max_row, 1, 9)
    provenance.freeze_panes = "A2"
    for col, width in enumerate([22, 20, 85, 110, 110, 85, 18, 18, 62], 1):
        provenance.column_dimensions[get_column_letter(col)].width = width

    paper = wb.create_sheet("PDF原TABLE3")
    paper.merge_cells("A1:C1")
    paper["A1"] = "PDF TABLE III 原始数值 — RMS-TRE（mm，mean ± std）"
    paper["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    paper["A1"].fill = PatternFill("solid", fgColor="7030A0")
    paper["A1"].alignment = Alignment(horizontal="center")
    paper.merge_cells("A2:C2")
    paper["A2"] = f"来源：{PDF_PATH}；印刷页 466（PDF 第 8 页）；visibility ∈ [0.2,0.3)；仅作论文原值对照"
    paper["A2"].alignment = Alignment(horizontal="center", wrap_text=True)
    for col, value in enumerate(["Method", "[0, 6) mm", "[6, 12] mm"], 1):
        paper.cell(4, col, value)
    style_header(paper, 4, 1, 3, color="7030A0")
    for method_name, values in PAPER_TABLE.items():
        paper.append([method_name] + [f"{mean:.2f} ± {std:.2f}" for mean, std in values])
    style_grid(paper, 4, paper.max_row, 1, 3)
    paper.column_dimensions["A"].width = 42
    paper.column_dimensions["B"].width = 22
    paper.column_dimensions["C"].width = 22

    readme = wb.create_sheet("说明")
    notes = [
        ["项目", "内容"],
        ["验证状态", "ANALYZED：仅汇总现有 JSON 结果，没有重新运行任何模型。"],
        ["指标", "RMS-TRE on paired volumetric fiducials，单位 mm；均值 ± 总体标准差（ddof=0）。"],
        ["筛选", "严格采用 0.2 ≤ visibility < 0.3；visibility=0.3 不纳入。"],
        ["形变分组", "[0,6) 与 [6,12] mm；6 mm 归入第二组，12 mm 纳入第二组，超过 12 mm 不纳入 TABLE III。"],
        ["数据条件", "使用 in silico、noise=0 mm 的现有评测结果。"],
        ["Lepard 元数据", "Lepard 与 Lepard+P2P 结果缺少 deformation_mm；按样本名与 Ours 同一数据集结果逐一匹配补充，匹配缺失为 0。"],
        ["Lepard 协议", "Lepard 当前为 RANSAC 基线；Lepard+P2P 为简化实现，二者均不能视为论文原实现复现。"],
        ["Ours 路径", "结果位于 RTORv8，但 JSON 内 checkpoint 指向 RTORv5 同名实验；保留迁移历史复核提示。"],
        ["PDF 原值", "PDF 原 TABLE III 单独放在 PDF原TABLE3 工作表，不用于填补或替换当前模型结果。"],
        ["统计谬误扫描", "11/11 已检查：本表为描述性分组汇总，不进行显著性、因果或个体层级推断；仍需注意模型间未配对检验及多重比较问题。"],
    ]
    for row in notes:
        readme.append(row)
    style_header(readme, 1, 1, 2, color="548235")
    style_grid(readme, 1, readme.max_row, 1, 2)
    readme.column_dimensions["A"].width = 22
    readme.column_dimensions["B"].width = 125

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    reopened = load_workbook(output, read_only=True)
    expected_sheets = ["TABLE3_当前结果", "数值明细", "数据来源", "PDF原TABLE3", "说明"]
    if reopened.sheetnames != expected_sheets:
        raise RuntimeError(f"Unexpected workbook sheets: {reopened.sheetnames}")
    if reopened["TABLE3_当前结果"].max_row != 4 + len(extracted):
        raise RuntimeError("Unexpected number of method rows")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output",
        type=Path,
        default=PROJECT / "results/paper_tables/TABLE3_in_silico_deformation_visibility_0.2_0.3_RMS_TRE_20260924.xlsx",
    )
    args = parser.parse_args()
    build_workbook(args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
