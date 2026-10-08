import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts.build_table7_excel import build_workbook, summarize_visibility_bins


class Table7ExcelTest(unittest.TestCase):
    def test_visibility_bin_boundaries(self):
        samples = [
            {"visibility": 0.2, "rms_tre_mm": 1.0},
            {"visibility": 0.2999, "rms_tre_mm": 3.0},
            {"visibility": 0.3, "rms_tre_mm": 5.0},
            {"visibility": 0.45, "rms_tre_mm": 5.0},
            {"visibility": 0.55, "rms_tre_mm": 5.0},
            {"visibility": 0.65, "rms_tre_mm": 5.0},
            {"visibility": 0.75, "rms_tre_mm": 5.0},
            {"visibility": 0.85, "rms_tre_mm": 5.0},
            {"visibility": 0.9999, "rms_tre_mm": 7.0},
            {"visibility": 1.0, "rms_tre_mm": 9.0},
            {"visibility": 1.0001, "rms_tre_mm": 100.0},
        ]
        result = summarize_visibility_bins(samples)
        self.assertEqual(result["[0.2, 0.3)"]["count"], 2)
        self.assertEqual(result["[0.2, 0.3)"]["mean_rms_tre_mm"], 2.0)
        self.assertEqual(result["[0.3, 0.4)"]["count"], 1)
        self.assertEqual(result["[0.9, 1.0]"]["count"], 2)
        self.assertEqual(result["[0.9, 1.0]"]["mean_rms_tre_mm"], 8.0)

    def test_workbook_preserves_source_status_and_pdf_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            result_path = root / "result.json"
            samples = []
            for i in range(8):
                samples.append({"visibility": 0.25 + 0.1 * i, "rms_tre_mm": float(i + 1)})
            result_path.write_text(
                json.dumps({"checkpoint": "model.pth", "evaluated_samples": 8, "samples": samples}),
                encoding="utf-8",
            )
            output = root / "table7.xlsx"
            methods = [{
                "name": "Ours",
                "env": "test",
                "requested": Path("model.pth"),
                "source_type": "json",
                "result": result_path,
                "validation": "已核对",
            }]
            build_workbook(output, methods=methods, low_visibility_reruns=[])
            wb = load_workbook(output, data_only=False)
            row = tuple(wb["TABLE7_当前汇总"].iter_rows(min_row=5, max_row=5, min_col=1, max_col=10, values_only=True))[0]
            self.assertEqual(row[0], "Ours")
            self.assertEqual(row[1], "1.00 ± 0.00")
            self.assertEqual(row[8], "8.00 ± 0.00")
            self.assertEqual(row[9], "已核对")
            self.assertEqual(wb["PDF原TABLE7"]["B10"].value, "14.97 ± 18.73")

    def test_low_visibility_rerun_accepts_two_bin_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            full_path = root / "full.json"
            rerun_path = root / "rerun.json"
            full_path.write_text(
                json.dumps({
                    "samples": [
                        {"visibility": 0.25 + 0.1 * i, "rms_tre_mm": float(i + 1)}
                        for i in range(8)
                    ]
                }),
                encoding="utf-8",
            )
            rerun_path.write_text(
                json.dumps({
                    "samples": [
                        {"visibility": 0.25, "rms_tre_mm": 2.0},
                        {"visibility": 0.35, "rms_tre_mm": 4.0},
                    ]
                }),
                encoding="utf-8",
            )
            output = root / "table7.xlsx"
            methods = [{
                "name": "Lepard",
                "env": "test",
                "requested": Path("model.pth"),
                "source_type": "json",
                "result": full_path,
                "validation": "已核对",
            }]
            reruns = [{"name": "Lepard", "result": rerun_path, "note": "two bins only"}]
            build_workbook(output, methods=methods, low_visibility_reruns=reruns)
            wb = load_workbook(output, data_only=False)
            row = tuple(wb["当前低可见率复跑"].iter_rows(min_row=2, max_row=2, min_col=1, max_col=11, values_only=True))[0]
            self.assertEqual(row[1], "2.00 ± 0.00")
            self.assertEqual(row[5], "4.00 ± 0.00")


if __name__ == "__main__":
    unittest.main()
