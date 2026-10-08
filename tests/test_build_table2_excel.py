import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts.build_table2_excel import build_workbook, summarize_low_visibility


class Table2ExcelTest(unittest.TestCase):
    def test_low_visibility_summary_excludes_upper_boundary(self):
        samples = [
            {"visibility": 0.2, "rms_tre_mm": 1.0},
            {"visibility": 0.2999, "rms_tre_mm": 3.0},
            {"visibility": 0.3, "rms_tre_mm": 100.0},
        ]

        summary = summarize_low_visibility(samples)

        self.assertEqual(summary["count"], 2)
        self.assertEqual(summary["mean_rms_tre_mm"], 2.0)
        self.assertEqual(summary["std_rms_tre_mm"], 1.0)

    def test_workbook_marks_missing_noise_and_flags_unverified_p2p(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            noise0 = root / "noise0.json"
            noise0.write_text(
                json.dumps(
                    {
                        "checkpoint": "/model.pth",
                        "evaluated_samples": 3,
                        "samples": [
                            {"visibility": 0.21, "rms_tre_mm": 2.0},
                            {"visibility": 0.25, "rms_tre_mm": 4.0},
                            {"visibility": 0.31, "rms_tre_mm": 50.0},
                        ],
                    }
                ),
                encoding="utf-8",
            )
            output = root / "table2.xlsx"
            methods = [
                {
                    "name": "Lepard+P2P",
                    "env": "livermatch5090",
                    "requested": Path("/model.pth"),
                    "results": {0: noise0, 2: None, 4: None},
                    "validation": "实现待复核",
                }
            ]

            build_workbook(output, methods=methods)

            workbook = load_workbook(output, read_only=True)
            rows = list(
                workbook["TABLE2_当前结果"].iter_rows(
                    min_row=5, max_row=5, values_only=True
                )
            )
            self.assertEqual(
                rows[0],
                ("Lepard+P2P", "3.00 ± 1.00", "未找到结果", "未找到结果", "1/3；实现待复核"),
            )
            self.assertEqual(
                workbook["PDF原TABLE2"]["B7"].value,
                "8.74 ± 10.23",
            )
            workbook.close()


if __name__ == "__main__":
    unittest.main()
