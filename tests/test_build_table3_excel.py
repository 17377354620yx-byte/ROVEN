import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts.build_table3_excel import build_workbook, summarize_deformation_bins


class Table3ExcelTest(unittest.TestCase):
    def test_deformation_and_visibility_boundaries(self):
        samples = [
            {"sample": "a", "visibility": 0.2, "deformation_mm": 0.0, "rms_tre_mm": 1.0},
            {"sample": "b", "visibility": 0.2999, "deformation_mm": 5.999, "rms_tre_mm": 3.0},
            {"sample": "c", "visibility": 0.25, "deformation_mm": 6.0, "rms_tre_mm": 5.0},
            {"sample": "d", "visibility": 0.25, "deformation_mm": 12.0, "rms_tre_mm": 7.0},
            {"sample": "e", "visibility": 0.3, "deformation_mm": 2.0, "rms_tre_mm": 100.0},
            {"sample": "f", "visibility": 0.25, "deformation_mm": 12.001, "rms_tre_mm": 100.0},
        ]
        result = summarize_deformation_bins(samples)
        self.assertEqual(result["[0, 6)"]["count"], 2)
        self.assertEqual(result["[0, 6)"]["mean_rms_tre_mm"], 2.0)
        self.assertEqual(result["[0, 6)"]["std_rms_tre_mm"], 1.0)
        self.assertEqual(result["[6, 12]"]["count"], 2)
        self.assertEqual(result["[6, 12]"]["mean_rms_tre_mm"], 6.0)
        self.assertEqual(result["[6, 12]"]["std_rms_tre_mm"], 1.0)

    def test_metadata_join_and_workbook_warning(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            result_path = root / "result.json"
            metadata_path = root / "metadata.json"
            result_path.write_text(
                json.dumps(
                    {
                        "checkpoint": "model.pth",
                        "evaluated_samples": 2,
                        "samples": [
                            {"sample_name": "a.npz", "visibility": 0.21, "RMS_TRE_mm": 3.0},
                            {"sample_name": "b.npz", "visibility": 0.22, "RMS_TRE_mm": 5.0},
                        ],
                    }
                ),
                encoding="utf-8",
            )
            metadata_path.write_text(
                json.dumps(
                    {
                        "samples": [
                            {"sample": "a.npz", "deformation_mm": 2.0},
                            {"sample": "b.npz", "deformation_mm": 8.0},
                        ]
                    }
                ),
                encoding="utf-8",
            )
            output = root / "table3.xlsx"
            methods = [
                {
                    "name": "Lepard+P2P",
                    "env": "test",
                    "requested": Path("model.pth"),
                    "result": result_path,
                    "metadata": metadata_path,
                    "validation": "简化 P2P 实现，待复核",
                }
            ]
            build_workbook(output, methods=methods)
            wb = load_workbook(output, data_only=False)
            self.assertEqual(
                tuple(wb["TABLE3_当前结果"].iter_rows(min_row=5, max_row=5, min_col=1, max_col=4, values_only=True))[0],
                ("Lepard+P2P", "3.00 ± 0.00", "5.00 ± 0.00", "实现待复核；deformation 元数据由样本名匹配补充"),
            )
            self.assertEqual(wb["PDF原TABLE3"]["B7"].value, "8.54 ± 11.16")


if __name__ == "__main__":
    unittest.main()
